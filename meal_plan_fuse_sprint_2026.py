import openpyxl
from openpyxl.styles import Font, Alignment, Border, Side, PatternFill
from openpyxl.utils import get_column_letter

wb = openpyxl.Workbook()
ws = wb.active
ws.title = "布勢スプリント2026"

thin = Side(style='thin')
border = Border(left=thin, right=thin, top=thin, bottom=thin)

def styled(ws, row, col, value, font=None, fill=None, align=None):
    c = ws.cell(row=row, column=col, value=value)
    if font: c.font = font
    if fill: c.fill = fill
    if align: c.alignment = align
    c.border = border
    return c

def borders(ws, r1, r2, c1, c2):
    for r in range(r1, r2+1):
        for c in range(c1, c2+1):
            ws.cell(row=r, column=c).border = border

def fill_range(ws, r1, r2, c1, c2, fill):
    for r in range(r1, r2+1):
        for c in range(c1, c2+1):
            ws.cell(row=r, column=c).fill = fill
            ws.cell(row=r, column=c).border = border

hdr_font = Font(name='Yu Gothic', bold=True, size=11)
norm_font = Font(name='Yu Gothic', size=10)
sm_font = Font(name='Yu Gothic', size=9)
sm_bold = Font(name='Yu Gothic', size=9, bold=True)
title_font = Font(name='Yu Gothic', bold=True, size=14)
note_font = Font(name='Yu Gothic', size=9, color='FF0000')
blue_note = Font(name='Yu Gothic', size=9, bold=True, color='0000AA')
center = Alignment(horizontal='center', vertical='center', wrap_text=True)
wrap_top = Alignment(wrap_text=True, vertical='top')
right_center = Alignment(horizontal='right', vertical='center')

green = PatternFill('solid', fgColor='C6EFCE')
blue = PatternFill('solid', fgColor='D6E4F0')
yellow = PatternFill('solid', fgColor='FFF2CC')
pink = PatternFill('solid', fgColor='FCE4EC')
orange = PatternFill('solid', fgColor='FBE5D6')
hdr_green = PatternFill('solid', fgColor='A9D18E')
gray = PatternFill('solid', fgColor='F2F2F2')
red_bg = PatternFill('solid', fgColor='FF6B6B')
orange_bg = PatternFill('solid', fgColor='FFB366')

# Column widths
ws.column_dimensions['A'].width = 10
for c in range(2, 8):
    ws.column_dimensions[get_column_letter(c)].width = 30

# === TITLE ===
ws.merge_cells('A1:G1')
styled(ws, 1, 1, '布勢スプリント2026　試合食事プラン', title_font, align=center)
ws.row_dimensions[1].height = 30

ws.merge_cells('A2:G2')
styled(ws, 2, 1, '100m（1本目 9:40 ／ 2本目 12:40）　体重64kg　目標糖質量 320〜384g/日（5〜6g/kg）', norm_font, align=center)
ws.row_dimensions[2].height = 22

# =============================================================
#  前日 (Row 4 ~ Row 15)
# =============================================================
R = 4  # start row

# Day label merged vertically
ws.merge_cells(start_row=R, start_column=1, end_row=R+11, end_column=1)
styled(ws, R, 1, "前日", Font(name='Yu Gothic', bold=True, size=13), align=center)
borders(ws, R, R+11, 1, 1)

meal_cols_before = ["朝食", "昼食", "間食", "夕食"]
for i, name in enumerate(meal_cols_before):
    styled(ws, R, 2+i, name, hdr_font, hdr_green, center)

# Row category labels (rows R+1 to R+7)
categories = [
    (R+1, "写真", yellow, 2),
    (R+3, "主食", green, 1),
    (R+4, "主菜", pink, 1),
    (R+5, "副菜", blue, 1),
    (R+6, "乳製品\n・果物", orange, 1),
    (R+7, "その他", gray, 1),
]

for row, label, fill_c, span in categories:
    if span > 1:
        # Don't merge col 1 as it's already merged for day label
        pass
    # These labels go in column 1 but col 1 is merged; we'll use a side approach
    # Actually, let me restructure: use col A for day, col B onward for meals
    # But the original format has the label column inside the day block
    pass

# Let me restructure: Col A = day label, Col B = row category, Col C-F = meals
# This better matches the original format

wb2 = openpyxl.Workbook()
ws = wb2.active
ws.title = "布勢スプリント2026"

ws.column_dimensions['A'].width = 6
ws.column_dimensions['B'].width = 9
for c_letter in ['C', 'D', 'E', 'F', 'G', 'H']:
    ws.column_dimensions[c_letter] = ws.column_dimensions['C']
    ws.column_dimensions[c_letter].width = 28

# === TITLE ===
ws.merge_cells('A1:G1')
styled(ws, 1, 1, '布勢スプリント2026　試合食事プラン', title_font, align=center)
ws.row_dimensions[1].height = 30

ws.merge_cells('A2:G2')
styled(ws, 2, 1, '100m（1本目 9:40 ／ 2本目 12:40）　体重64kg　目標糖質量 320〜384g/日（5〜6g/kg）', norm_font, align=center)

# =============================================================
# 前日
# =============================================================
R = 4
num_rows = 10  # total rows for this block

# Col A: Day label
ws.merge_cells(start_row=R, start_column=1, end_row=R+num_rows-1, end_column=1)
styled(ws, R, 1, "前日", Font(name='Yu Gothic', bold=True, size=13), align=center)
borders(ws, R, R+num_rows-1, 1, 1)

# Col B: Category labels
cat_rows = [
    (R, ""),       # header row
    (R+1, "写真"),
    (R+2, "主食"),
    (R+3, "主菜"),
    (R+4, "副菜"),
    (R+5, "乳製品\n・果物"),
    (R+6, "その他"),
]
cat_fills = [hdr_green, yellow, green, pink, blue, orange, gray]
for (row, label), fl in zip(cat_rows, cat_fills):
    styled(ws, row, 2, label, sm_bold, fl, center)

# Content label spans rows R+7 to R+9
ws.merge_cells(start_row=R+7, start_column=2, end_row=R+9, end_column=2)
styled(ws, R+7, 2, "内容", sm_bold, align=center)
borders(ws, R+7, R+9, 2, 2)

# Meal headers (Col C-F)
meals_before = ["朝食", "昼食", "間食", "夕食"]
for i, m in enumerate(meals_before):
    styled(ws, R, 3+i, m, hdr_font, hdr_green, center)

# Category row fills (C-F)
for i in range(4):
    col = 3 + i
    styled(ws, R+1, col, "", sm_font, yellow, center)   # 写真
    styled(ws, R+2, col, "", sm_font, green, center)     # 主食
    styled(ws, R+3, col, "", sm_font, pink, center)      # 主菜
    styled(ws, R+4, col, "", sm_font, blue, center)      # 副菜
    styled(ws, R+5, col, "", sm_font, orange, center)    # 乳製品
    styled(ws, R+6, col, "", sm_font, gray, center)      # その他

# Content data
before_content = [
    # 朝食
    [
        "白米 250g（糖質≒92g）",
        "サケ 1切れ",
        "味噌汁（わかめ、豆腐）",
        "サラダ（レタス、トマト2個）",
        "ヨーグルト",
        "ブルーベリー 20粒くらい",
        "豆乳",
    ],
    # 昼食
    [
        "白米 300g（糖質≒110g）",
        "豚肉のポークチャップ",
        "サラダ（たまご1個、レタス、トマト2個）",
        "あおさの味噌汁",
        "パンプキンポテトサラダ",
        "ヨーグルト",
        "フルーツ（オレンジ、キウイ）",
    ],
    # 間食
    [
        "焼き芋 1本（糖質≒50g）",
        "アイスコーヒー",
        "クレアチン 5g",
    ],
    # 夕食
    [
        "白米 250g（糖質≒92g）",
        "生姜焼き（豚肉、玉ねぎ）",
        "ブロッコリー",
        "サラダ（たまご1個、トマト2個、レタス、鶏肉）",
        "ヨーグルト",
        "ゴールドキウイ",
    ],
]

for i, items in enumerate(before_content):
    col = 3 + i
    ws.merge_cells(start_row=R+7, start_column=col, end_row=R+9, end_column=col)
    styled(ws, R+7, col, "\n".join(items), sm_font, align=wrap_top)
    borders(ws, R+7, R+9, col, col)

# Row heights for Day Before
ws.row_dimensions[R+1].height = 50    # 写真
for rr in range(R+2, R+7):
    ws.row_dimensions[rr].height = 20
for rr in range(R+7, R+10):
    ws.row_dimensions[rr].height = 40

# Summary
SR = R + num_rows
ws.merge_cells(start_row=SR, start_column=1, end_row=SR, end_column=6)
styled(ws, SR, 1,
       '【前日の糖質目安】白米250g×2＋300g×1＋焼き芋≒344g ＋果物・野菜等 → 約360〜380g（5.6〜5.9g/kg）✓',
       blue_note, align=Alignment(wrap_text=True))

SR2 = SR + 1
ws.merge_cells(start_row=SR2, start_column=1, end_row=SR2, end_column=6)
styled(ws, SR2, 1,
       '※前日は脂質を控えめにし、消化の良い糖質中心の食事で筋グリコーゲンを蓄える。食物繊維の多い生野菜は少なめに。',
       note_font, align=Alignment(wrap_text=True))

# =============================================================
# 当日
# =============================================================
R2 = SR2 + 2
num_rows2 = 10

ws.merge_cells(start_row=R2, start_column=1, end_row=R2+num_rows2-1, end_column=1)
styled(ws, R2, 1, "当日", Font(name='Yu Gothic', bold=True, size=13), align=center)
borders(ws, R2, R2+num_rows2-1, 1, 1)

# Category labels
cat_rows2 = [
    (R2, ""),
    (R2+1, "写真"),
    (R2+2, "主食"),
    (R2+3, "主菜"),
    (R2+4, "副菜"),
    (R2+5, "乳製品\n・果物"),
    (R2+6, "その他"),
]
for (row, label), fl in zip(cat_rows2, cat_fills):
    styled(ws, row, 2, label, sm_bold, fl, center)

ws.merge_cells(start_row=R2+7, start_column=2, end_row=R2+9, end_column=2)
styled(ws, R2+7, 2, "内容", sm_bold, align=center)
borders(ws, R2+7, R2+9, 2, 2)

# Competition day has 5 meal slots
comp_meals = [
    "朝食\n【6:00頃】\n(1本目3.5h前)",
    "補食①\n【8:30〜9:00】\n(1本目1h前)",
    "補食②\n【10:00〜10:30】\n(1本目後→2本目前)",
    "補食③\n【レース後】\n(2本目終了後)",
    "夕食\n【18:00〜】",
]

for i, m in enumerate(comp_meals):
    styled(ws, R2, 3+i, m, Font(name='Yu Gothic', bold=True, size=9), hdr_green, center)
ws.row_dimensions[R2].height = 50

# Category fills for 5 columns
for i in range(5):
    col = 3 + i
    styled(ws, R2+1, col, "", sm_font, yellow, center)
    styled(ws, R2+2, col, "", sm_font, green, center)
    styled(ws, R2+3, col, "", sm_font, pink, center)
    styled(ws, R2+4, col, "", sm_font, blue, center)
    styled(ws, R2+5, col, "", sm_font, orange, center)
    styled(ws, R2+6, col, "", sm_font, gray, center)

comp_content = [
    # 朝食
    [
        "白米 250g（糖質≒92g）",
        "サケ 1切れ",
        "味噌汁（豆腐、ネギ）",
        "たまご焼き 1個分",
        "ヨーグルト",
        "ゴールドキウイ",
    ],
    # 補食①
    [
        "バナナ 1本（糖質≒25g）",
        "エネルギーゼリー 1個（糖質≒25g）",
        "水 or スポーツドリンク",
    ],
    # 補食②
    [
        "おにぎり 1個（糖質≒40g）",
        "バナナ 1本（糖質≒25g）",
        "スポーツドリンク",
        "※招集前に消化を終えるよう",
        "　10:50頃までに食べ終える",
    ],
    # 補食③
    [
        "おにぎり 1個（糖質≒40g）",
        "オレンジジュース",
        "プロテイン",
        "クレアチン 5g",
    ],
    # 夕食
    [
        "白米 250g（糖質≒92g）",
        "鶏肉（グリルまたはソテー）",
        "サラダ（レタス、トマト2個、たまご1個）",
        "味噌汁",
        "ヨーグルト",
        "フルーツ（キウイ、オレンジ）",
    ],
]

for i, items in enumerate(comp_content):
    col = 3 + i
    ws.merge_cells(start_row=R2+7, start_column=col, end_row=R2+9, end_column=col)
    styled(ws, R2+7, col, "\n".join(items), sm_font, align=wrap_top)
    borders(ws, R2+7, R2+9, col, col)

ws.row_dimensions[R2+1].height = 50
for rr in range(R2+2, R2+7):
    ws.row_dimensions[rr].height = 20
for rr in range(R2+7, R2+10):
    ws.row_dimensions[rr].height = 45

# Summary
SR3 = R2 + num_rows2
ws.merge_cells(start_row=SR3, start_column=1, end_row=SR3, end_column=7)
styled(ws, SR3, 1,
       '【当日の糖質目安】白米250g×2＋おにぎり2個＋バナナ2本＋ゼリー＋果物等≒370〜390g（5.8〜6.1g/kg）✓',
       blue_note, align=Alignment(wrap_text=True))

SR4 = SR3 + 1
ws.merge_cells(start_row=SR4, start_column=1, end_row=SR4+1, end_column=7)
styled(ws, SR4, 1,
       '※当日ポイント\n'
       '・朝食は競技開始3〜4時間前までに済ませる（6:00頃）\n'
       '・補食①は1本目の60〜90分前に消化しやすいもの（バナナ、ゼリー）\n'
       '・1本目→2本目の間（約3時間）でおにぎり＋バナナで糖質補給。招集（12:20）の1.5時間前（10:50頃まで）に食べ終える\n'
       '・脂質の多いもの（揚げ物、チョコレート等）は当日レース前は避ける\n'
       '・水分はこまめに。スポーツドリンクで電解質も補給',
       sm_font, align=wrap_top)
ws.row_dimensions[SR4].height = 40
ws.row_dimensions[SR4+1].height = 50

# =============================================================
# タイムスケジュール
# =============================================================
TS = SR4 + 3
ws.merge_cells(start_row=TS, start_column=1, end_row=TS, end_column=5)
styled(ws, TS, 1, '【当日タイムスケジュール】', hdr_font, align=Alignment(vertical='center'))

schedule = [
    ("6:00", "起床・朝食", None),
    ("8:30〜9:00", "補食①（バナナ＋ゼリー）", green),
    ("9:00", "ウォームアップ開始", None),
    ("9:20", "招集完了（1本目）", orange_bg),
    ("9:40", "競技開始（1本目）", red_bg),
    ("10:00〜10:30", "補食②（おにぎり＋バナナ）", green),
    ("11:00〜", "ウォームアップ（2本目）", None),
    ("12:20", "招集完了（2本目）", orange_bg),
    ("12:40", "競技開始（2本目）", red_bg),
    ("13:00〜", "補食③（おにぎり＋OJ＋プロテイン＋クレアチン）", green),
    ("18:00〜", "夕食（リカバリー）", None),
]

for i, (time_str, action, bg) in enumerate(schedule):
    row = TS + 1 + i
    styled(ws, row, 1, time_str, Font(name='Yu Gothic', bold=True, size=10), bg, right_center)
    ws.merge_cells(start_row=row, start_column=2, end_row=row, end_column=5)
    styled(ws, row, 2, action, norm_font, bg, Alignment(vertical='center'))
    for c in range(2, 6):
        ws.cell(row=row, column=c).border = border
        if bg:
            ws.cell(row=row, column=c).fill = bg

# =============================================================
# サプリメント
# =============================================================
SP = TS + len(schedule) + 2
ws.merge_cells(start_row=SP, start_column=1, end_row=SP, end_column=5)
styled(ws, SP, 1, '【サプリメント】', hdr_font, align=Alignment(vertical='center'))

SP2 = SP + 1
ws.merge_cells(start_row=SP2, start_column=1, end_row=SP2+1, end_column=5)
styled(ws, SP2, 1,
       '・クレアチン 5g/日（前日：間食時、当日：レース後）\n'
       '・プロテイン（当日レース後のリカバリーに）',
       norm_font, align=wrap_top)

# Save
output_path = '/home/user/AK/布勢スプリント2026_食事プラン.xlsx'
wb2.save(output_path)
print(f"Saved: {output_path}")
